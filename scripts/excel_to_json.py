# -*- coding: utf-8 -*-
"""
Hoyeh Cafeteria Excel -> JSON
Compatible with existing menu-data-full.json structure

Outputs:
1. menu-data-full.json  -> all known dates
2. menu-data.json       -> rolling window (yesterday to +30 days)

Expected Excel structure (per sheet Week 1 ~ Week N):
- C2:I2 => Mon .. Sun
- C3:I3 => dates
- Column A contains BREAKFAST, LUNCH, DINNER, and SUPPER section labels
- Column B contains the row labels used to map each section

Meal row numbers are detected from the labels instead of being fixed. This
supports menu templates where, for example, Rice & Spice contains five rows in
one month and six rows in another month.
"""

import json
import os
import re
import sys
from copy import deepcopy
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


DATA_DIR = "data"

FULL_JSON_FILE = os.path.join(DATA_DIR, "menu-data-full.json")
ROLLING_JSON_FILE = os.path.join(DATA_DIR, "menu-data.json")

# Rolling window: yesterday to future 30 days
ROLLING_DAYS_BEFORE = 1
ROLLING_DAYS_AFTER = 30

EXPECTED_WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

DATE_HEADER_ROW = 3
WEEKDAY_HEADER_ROW = 2
START_COL = 3   # C
END_COL = 9     # I

MEAL_LABELS = ("BREAKFAST", "LUNCH", "DINNER", "SUPPER")

MAIN_MEAL_ROW_LABELS = (
    "The Bowl",
    "The Culinary Table",
    "Rice & Spice",
    "Green House",
    "Soup",
    "Dessert",
)


class ValidationError(Exception):
    pass


def log_info(msg: str):
    print(f"[INFO] {msg}")


def log_warning(msg: str):
    print(f"[WARNING] {msg}")


def log_error(msg: str):
    print(f"[ERROR] {msg}")


def read_json_file(path: str) -> dict:
    if not os.path.exists(path):
        log_info(f"{path} not found. Starting with empty dataset.")
        return {}

    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            raise ValidationError(f"{path} must contain a JSON object.")
        log_info(f"Loaded existing data from {path} ({len(data)} date(s))")
        return data
    except json.JSONDecodeError as e:
        raise ValidationError(f"Failed to parse JSON file '{path}': {e}")
    except Exception as e:
        raise ValidationError(f"Failed to read JSON file '{path}': {e}")


def save_json_file(path: str, data: dict):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    log_info(f"Saved {path} ({len(data)} date(s))")


def excel_col_letter(col_idx: int) -> str:
    result = ""
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result


def normalize_text(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    # normalize multiple spaces inside each line, but preserve line breaks
    lines = []
    for line in text.split("\n"):
        cleaned = re.sub(r"[ \t]+", " ", line).strip()
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def normalize_label(value) -> str:
    """Normalize structural labels while keeping menu item text untouched."""
    return " ".join(normalize_text(value).split()).casefold()


def split_lines(value) -> list[str]:
    """
    Split cell text into lines.
    Used for Vegetarian or cells with internal line breaks.
    Keeps the order, strips numbering like '1.' / '2)' / '3 ' from the beginning.
    """
    text = normalize_text(value)
    if not text:
        return []

    items = []
    for line in text.split("\n"):
        cleaned = re.sub(r"^\s*\d+\s*[\.\)]\s*", "", line).strip()
        cleaned = re.sub(r"^\s*\d+\s+", "", cleaned).strip()
        if cleaned:
            items.append(cleaned)
    return items


def get_cell_text(ws, cell_ref: str) -> str:
    return normalize_text(ws[cell_ref].value)


def get_row_values(ws, row_num: int, col_num: int, warnings: list[str], context: str) -> list[str]:
    """
    Returns non-empty values from a single row + one column position.
    Since each date is one column, this reads only one cell.
    If the cell contains line breaks, preserve as one string for normal menu cells.
    """
    cell_ref = f"{excel_col_letter(col_num)}{row_num}"
    value = normalize_text(ws[cell_ref].value)
    if not value:
        warnings.append(f"{ws.title} {cell_ref}: blank cell in {context}")
        return []
    return [value]


def get_vertical_block(ws, row_start: int, row_end: int, col_num: int, warnings: list[str], context: str) -> list[str]:
    values = []
    for row_num in range(row_start, row_end + 1):
        cell_ref = f"{excel_col_letter(col_num)}{row_num}"
        value = normalize_text(ws[cell_ref].value)
        if value:
            values.append(value)
        else:
            warnings.append(f"{ws.title} {cell_ref}: blank cell in {context}")
    return values


def parse_excel_date(value, sheet_name: str, cell_ref: str) -> date:
    if value is None:
        raise ValidationError(f"Sheet '{sheet_name}' cell {cell_ref}: date is blank.")

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()
    if not text:
        raise ValidationError(f"Sheet '{sheet_name}' cell {cell_ref}: date is blank.")

    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    raise ValidationError(f"Sheet '{sheet_name}' cell {cell_ref}: invalid date '{value}'.")


def format_date_text(dt: date) -> str:
    return dt.strftime("%d/%m/%Y")


def format_date_key(dt: date) -> str:
    return dt.strftime("%Y%m%d")


def normalize_repo_path(path_text: str) -> str:
    normalized = path_text.replace("\\", "/").strip()
    while normalized.startswith("./"):
        normalized = normalized[2:]
    return normalized


def is_source_excel_path(path_text: str) -> bool:
    normalized = normalize_repo_path(path_text)
    filename = Path(normalized).name
    return (
        normalized.startswith("source/")
        and normalized.lower().endswith(".xlsx")
        and not filename.startswith("~$")
    )


def existing_source_excel_path(path_text: str) -> Path | None:
    if not is_source_excel_path(path_text):
        return None

    path = Path(normalize_repo_path(path_text))
    if path.exists() and path.is_file():
        return path

    return None


def find_excel_files_from_github_event() -> list[str]:
    """
    In GitHub Actions, checkout file mtimes are unreliable. Prefer the workbook
    files that actually appeared in the push event so newly uploaded or modified
    menus win, including older-month corrections.

    Every changed workbook is returned in event order. Processing only the file
    with the latest menu date would silently skip an older-month correction when
    several workbooks are uploaded in the same push.
    """
    event_path = os.environ.get("GITHUB_EVENT_PATH", "").strip()
    if not event_path:
        return []

    try:
        with open(event_path, "r", encoding="utf-8") as f:
            event = json.load(f)
    except Exception as e:
        log_warning(f"Unable to read GitHub event payload '{event_path}': {e}")
        return []

    changed_candidates = []

    commits = event.get("commits") or []
    if not commits and isinstance(event.get("head_commit"), dict):
        commits = [event["head_commit"]]

    for commit in commits:
        if not isinstance(commit, dict):
            continue

        for change_type in ("added", "modified"):
            for path_text in commit.get(change_type, []) or []:
                if is_source_excel_path(path_text):
                    changed_candidates.append(path_text)

    # Resolve existing paths and de-duplicate while preserving event order.
    existing_candidates = {}

    for path_text in changed_candidates:
        path = existing_source_excel_path(path_text)
        if path is None:
            continue

        normalized = normalize_repo_path(str(path))
        # If one path changed in several commits in the push, process only the
        # final checked-out version and keep it at its latest event position.
        if normalized in existing_candidates:
            del existing_candidates[normalized]
        existing_candidates[normalized] = path

    if existing_candidates:
        selected = [str(path) for path in existing_candidates.values()]
        log_info(
            "Using Excel file(s) from GitHub push event: "
            + ", ".join(selected)
        )
        return selected

    if changed_candidates:
        names = ", ".join(normalize_repo_path(p) for p in changed_candidates)
        log_warning(f"GitHub event referenced Excel file(s) not found after checkout: {names}")

    return []


def read_workbook_date_span(excel_path: Path) -> tuple[date, date] | None:
    try:
        wb = load_workbook(excel_path, data_only=True, read_only=True)
    except Exception as e:
        log_warning(f"Unable to inspect dates in '{excel_path}': {e}")
        return None

    try:
        dates = []
        week_sheets = [name for name in wb.sheetnames if name.strip().lower().startswith("week")]
        for sheet_name in week_sheets:
            ws = wb[sheet_name]
            for col in range(START_COL, END_COL + 1):
                cell_ref = f"{excel_col_letter(col)}{DATE_HEADER_ROW}"
                value = ws.cell(DATE_HEADER_ROW, col).value
                if value is None:
                    continue
                try:
                    dates.append(parse_excel_date(value, sheet_name, cell_ref))
                except ValidationError:
                    continue

        if not dates:
            return None

        return min(dates), max(dates)
    finally:
        wb.close()


def find_latest_source_excel_file(source_candidates: list[Path]) -> str:
    dated_candidates = []

    for path in source_candidates:
        span = read_workbook_date_span(path)
        if span is not None:
            first_date, last_date = span
            dated_candidates.append((last_date, first_date, path.name.lower(), path))

    if dated_candidates:
        last_date, _, _, latest_file = max(dated_candidates)
        log_info(
            f"Using Excel file from source/ with latest menu date "
            f"{format_date_text(last_date)}: {latest_file}"
        )
        return str(latest_file)

    names = ", ".join(sorted(p.name for p in source_candidates))
    raise ValidationError(
        "Could not determine menu dates from any source/ workbook. "
        "Refusing to select by filesystem modified time because GitHub Actions "
        f"checkout mtimes are unreliable. Candidate files: {names}"
    )


def find_unique_label_row(
    ws,
    col_num: int,
    label: str,
    row_start: int = 1,
    row_end: int | None = None,
) -> int:
    """Find exactly one normalized label in a worksheet column and row range."""
    row_end = row_end or ws.max_row
    expected = normalize_label(label)
    matches = [
        row_num
        for row_num in range(row_start, row_end + 1)
        if normalize_label(ws.cell(row_num, col_num).value) == expected
    ]

    col_letter = excel_col_letter(col_num)
    location = f"{col_letter}{row_start}:{col_letter}{row_end}"
    if not matches:
        raise ValidationError(
            f"Sheet '{ws.title}' range {location}: required label '{label}' was not found."
        )
    if len(matches) > 1:
        cells = ", ".join(f"{col_letter}{row_num}" for row_num in matches)
        raise ValidationError(
            f"Sheet '{ws.title}': label '{label}' must appear once in {location}; found at {cells}."
        )

    return matches[0]


def detect_sheet_layout(ws) -> dict:
    """
    Detect meal boundaries from column A and row roles from column B.

    Column A may contain merged cells, so only the first row of each meal block
    carries a value. Column B is used to identify the variable-length rows inside
    lunch, dinner, and supper. The Vegetarian row is intentionally found in
    column B because some older workbooks also contain a stray Vegetarian value
    in column A on the supper Dessert row.
    """
    meal_rows = {
        label.lower(): find_unique_label_row(ws, 1, label)
        for label in MEAL_LABELS
    }

    ordered_meal_rows = [meal_rows[label.lower()] for label in MEAL_LABELS]
    if ordered_meal_rows != sorted(ordered_meal_rows) or len(set(ordered_meal_rows)) != len(ordered_meal_rows):
        positions = ", ".join(
            f"{label}=A{meal_rows[label.lower()]}" for label in MEAL_LABELS
        )
        raise ValidationError(
            f"Sheet '{ws.title}': meal sections must appear in the order "
            f"BREAKFAST, LUNCH, DINNER, SUPPER; found {positions}."
        )

    vegetarian_row = find_unique_label_row(
        ws,
        2,
        "Vegetarian",
        row_start=meal_rows["supper"] + 1,
    )

    raw_ranges = {
        "breakfast": (meal_rows["breakfast"], meal_rows["lunch"] - 1),
        "lunch": (meal_rows["lunch"], meal_rows["dinner"] - 1),
        "dinner": (meal_rows["dinner"], meal_rows["supper"] - 1),
        "supper": (meal_rows["supper"], vegetarian_row - 1),
    }

    for section, (start_row, end_row) in raw_ranges.items():
        if start_row > end_row:
            raise ValidationError(
                f"Sheet '{ws.title}': invalid {section.upper()} range {start_row}..{end_row}."
            )

    layout = {
        "breakfast": {
            "start": raw_ranges["breakfast"][0],
            "end": raw_ranges["breakfast"][1],
        },
        "vegetarian_row": vegetarian_row,
    }

    for section in ("lunch", "dinner", "supper"):
        start_row, end_row = raw_ranges[section]
        category_rows = {
            label: find_unique_label_row(
                ws,
                2,
                label,
                row_start=start_row,
                row_end=end_row,
            )
            for label in MAIN_MEAL_ROW_LABELS
        }

        ordered_category_rows = [category_rows[label] for label in MAIN_MEAL_ROW_LABELS]
        if ordered_category_rows != sorted(ordered_category_rows):
            positions = ", ".join(
                f"{label}=B{category_rows[label]}" for label in MAIN_MEAL_ROW_LABELS
            )
            raise ValidationError(
                f"Sheet '{ws.title}' {section.upper()}: row labels are out of order; {positions}."
            )

        if category_rows["The Bowl"] != start_row:
            raise ValidationError(
                f"Sheet '{ws.title}' {section.upper()}: expected 'The Bowl' at B{start_row}, "
                f"found it at B{category_rows['The Bowl']}."
            )

        if category_rows["Dessert"] != end_row:
            raise ValidationError(
                f"Sheet '{ws.title}' {section.upper()}: expected 'Dessert' at the final row "
                f"B{end_row}, found it at B{category_rows['Dessert']}."
            )

        layout[section] = {
            "start": start_row,
            "end": end_row,
            "rows": category_rows,
        }

    return layout


def validate_sheet_structure(ws) -> dict:
    layout = detect_sheet_layout(ws)

    # Weekday headers
    for idx, expected in enumerate(EXPECTED_WEEKDAYS, start=START_COL):
        cell_ref = f"{excel_col_letter(idx)}{WEEKDAY_HEADER_ROW}"
        actual = get_cell_text(ws, cell_ref)
        if actual != expected:
            raise ValidationError(
                f"Sheet '{ws.title}' cell {cell_ref}: expected weekday '{expected}', got '{actual or '[blank]'}'."
            )

    # Dates present and continuous
    dates = []
    for col in range(START_COL, END_COL + 1):
        cell_ref = f"{excel_col_letter(col)}{DATE_HEADER_ROW}"
        dt = parse_excel_date(ws[cell_ref].value, ws.title, cell_ref)
        dates.append(dt)

    for i in range(1, len(dates)):
        if dates[i] != dates[i - 1] + timedelta(days=1):
            prev_ref = f"{excel_col_letter(START_COL + i - 1)}{DATE_HEADER_ROW}"
            curr_ref = f"{excel_col_letter(START_COL + i)}{DATE_HEADER_ROW}"
            raise ValidationError(
                f"Sheet '{ws.title}': date sequence broken between {prev_ref} ({format_date_text(dates[i - 1])}) "
                f"and {curr_ref} ({format_date_text(dates[i])})."
            )

    return layout


def build_breakfast(ws, col: int, dt: date, warnings: list[str], section: dict) -> dict:
    lines = get_vertical_block(
        ws, section["start"], section["end"], col, warnings,
        context=f"BREAKFAST {format_date_text(dt)}"
    )
    if not lines:
        log_warning(f"Sheet '{ws.title}' date {format_date_text(dt)}: breakfast is empty.")

    return {
        "date_text": format_date_text(dt),
        "lines": lines,
        "title": "BREAKFAST MENU",
    }


def build_lunch(ws, col: int, dt: date, warnings: list[str], section: dict) -> dict:
    rows = section["rows"]
    noodle = get_row_values(ws, rows["The Bowl"], col, warnings, f"LUNCH noodleSpecial {format_date_text(dt)}")
    daily = []
    for label in ("The Culinary Table", "Green House", "Soup", "Dessert"):
        daily.extend(get_row_values(ws, rows[label], col, warnings, f"LUNCH dailySpecial {format_date_text(dt)}"))
    home = get_vertical_block(
        ws,
        rows["Rice & Spice"],
        rows["Green House"] - 1,
        col,
        warnings,
        f"LUNCH homeTaste {format_date_text(dt)}",
    )

    return {
        "dailySpecial": daily,
        "dailySpecialTitle": "DAILY SPECIAL",
        "date_text": format_date_text(dt),
        "homeTaste": home,
        "homeTasteTitle": "HOMETASTE DAILY SELECTION",
        "noodleSpecial": noodle,
        "noodleSpecialTitle": "NOODLE SPECIAL",
        "title": "LUNCH MENU",
    }


def build_dinner(ws, col: int, dt: date, warnings: list[str], section: dict) -> dict:
    rows = section["rows"]
    noodle = get_row_values(ws, rows["The Bowl"], col, warnings, f"DINNER noodleSpecial {format_date_text(dt)}")
    daily = []
    for label in ("The Culinary Table", "Green House", "Soup", "Dessert"):
        daily.extend(get_row_values(ws, rows[label], col, warnings, f"DINNER dailySpecial {format_date_text(dt)}"))
    home = get_vertical_block(
        ws,
        rows["Rice & Spice"],
        rows["Green House"] - 1,
        col,
        warnings,
        f"DINNER homeTaste {format_date_text(dt)}",
    )

    return {
        "dailySpecial": daily,
        "dailySpecialTitle": "DAILY SPECIAL",
        "date_text": format_date_text(dt),
        "homeTaste": home,
        "homeTasteTitle": "HOMETASTE DAILY SELECTION",
        "noodleSpecial": noodle,
        "noodleSpecialTitle": "NOODLE SPECIAL",
        "title": "DINNER MENU",
    }


def build_supper(ws, col: int, dt: date, warnings: list[str], section: dict) -> dict:
    rows = section["rows"]
    noodle = get_row_values(ws, rows["The Bowl"], col, warnings, f"SUPPER noodleSpecial {format_date_text(dt)}")
    daily = []
    for label in ("The Culinary Table", "Soup", "Dessert"):
        daily.extend(get_row_values(ws, rows[label], col, warnings, f"SUPPER dailySpecial {format_date_text(dt)}"))
    home = get_vertical_block(
        ws,
        rows["Rice & Spice"],
        rows["Green House"],
        col,
        warnings,
        f"SUPPER homeTaste {format_date_text(dt)}",
    )

    return {
        "dailySpecial": daily,
        "dailySpecialTitle": "DAILY SPECIAL",
        "date_text": format_date_text(dt),
        "homeTaste": home,
        "homeTasteTitle": "HOMETASTE DAILY SELECTION",
        "noodleSpecial": noodle,
        "noodleSpecialTitle": "NOODLE SPECIAL",
        "title": "SUPPER MENU",
    }


def build_vegetarian(ws, col: int, dt: date, warnings: list[str], row_num: int) -> dict:
    cell_ref = f"{excel_col_letter(col)}{row_num}"
    raw_text = normalize_text(ws[cell_ref].value)
    lines = split_lines(ws[cell_ref].value)

    if not raw_text:
        warnings.append(f"{ws.title} {cell_ref}: Vegetarian menu is blank for {format_date_text(dt)}")

    return {
        "date_text": format_date_text(dt),
        "title": "VEGETARIAN MENU",
        "lines": lines,
        "raw_text": raw_text,
    }


def parse_sheet(ws, warnings: list[str]) -> dict:
    layout = validate_sheet_structure(ws)
    log_info(
        f"Detected layout for sheet '{ws.title}': "
        f"BREAKFAST {layout['breakfast']['start']}-{layout['breakfast']['end']}, "
        f"LUNCH {layout['lunch']['start']}-{layout['lunch']['end']}, "
        f"DINNER {layout['dinner']['start']}-{layout['dinner']['end']}, "
        f"SUPPER {layout['supper']['start']}-{layout['supper']['end']}, "
        f"Vegetarian {layout['vegetarian_row']}"
    )

    parsed = {}
    for col in range(START_COL, END_COL + 1):
        dt = parse_excel_date(ws.cell(DATE_HEADER_ROW, col).value, ws.title, f"{excel_col_letter(col)}{DATE_HEADER_ROW}")
        date_key = format_date_key(dt)

        if date_key in parsed:
            raise ValidationError(f"Sheet '{ws.title}': duplicate date '{date_key}' within sheet.")

        parsed[date_key] = {
            "breakfast": build_breakfast(ws, col, dt, warnings, layout["breakfast"]),
            "lunch": build_lunch(ws, col, dt, warnings, layout["lunch"]),
            "dinner": build_dinner(ws, col, dt, warnings, layout["dinner"]),
            "supper": build_supper(ws, col, dt, warnings, layout["supper"]),
            "vegetarian": build_vegetarian(ws, col, dt, warnings, layout["vegetarian_row"]),
        }

    return parsed


def find_excel_files() -> list[str]:
    """
    Priority:
    1. MENU_EXCEL_FILE env var
    2. all .xlsx files uploaded/modified in the current GitHub push event
    3. source/ .xlsx file with the latest menu dates
    4. menu.xlsx in repo root
    5. single .xlsx in repo root
    """
    # Priority 1: env var
    env_path = os.environ.get("MENU_EXCEL_FILE", "").strip()
    if env_path:
        if os.path.exists(env_path):
            return [env_path]
        raise ValidationError(f"MENU_EXCEL_FILE is set but file not found: {env_path}")

    # Priority 2: every workbook from the current GitHub push event
    github_event_files = find_excel_files_from_github_event()
    if github_event_files:
        return github_event_files

    # Priority 3: source/ workbook with the latest menu date
    source_dir = Path("source")
    if source_dir.exists() and source_dir.is_dir():
        source_candidates = [
            p for p in source_dir.glob("*.xlsx")
            if not p.name.startswith("~$")
        ]
        if source_candidates:
            return [find_latest_source_excel_file(source_candidates)]

    # Priority 4: menu.xlsx in repo root
    preferred = Path("menu.xlsx")
    if preferred.exists():
        return [str(preferred)]

    # Priority 5: single .xlsx file in repo root
    root_candidates = [
        p for p in Path(".").glob("*.xlsx")
        if not p.name.startswith("~$")
    ]
    if len(root_candidates) == 1:
        return [str(root_candidates[0])]

    if len(root_candidates) == 0:
        raise ValidationError(
            "No Excel file found. Put an .xlsx file into 'source/' folder, or set MENU_EXCEL_FILE."
        )

    names = ", ".join(sorted(p.name for p in root_candidates))
    raise ValidationError(
        f"Multiple Excel files found in repo root: {names}. "
        f"Please put the target Excel file into 'source/' folder, or set MENU_EXCEL_FILE."
    )


def load_excel_data(excel_path: str) -> tuple[dict, list[str]]:
    warnings = []
    try:
        wb = load_workbook(excel_path, data_only=True)
    except Exception as e:
        raise ValidationError(f"Failed to open Excel file '{excel_path}': {e}")

    if not wb.sheetnames:
        raise ValidationError(f"Excel file '{excel_path}' has no sheets.")

    week_sheets = [name for name in wb.sheetnames if name.strip().lower().startswith("week")]
    if not week_sheets:
        raise ValidationError(f"Excel file '{excel_path}' contains no 'Week' sheets.")

    log_info(f"Reading Excel file: {excel_path}")
    log_info(f"Week sheets found: {', '.join(week_sheets)}")

    all_data = {}

    for sheet_name in week_sheets:
        ws = wb[sheet_name]
        log_info(f"Parsing sheet: {sheet_name}")
        sheet_data = parse_sheet(ws, warnings)

        for date_key, payload in sheet_data.items():
            if date_key in all_data:
                raise ValidationError(
                    f"Duplicate date across sheets: {date_key} already exists before sheet '{sheet_name}'."
                )
            all_data[date_key] = payload

    return dict(sorted(all_data.items())), warnings


def merge_full_data(existing_full: dict, new_data: dict) -> dict:
    merged = deepcopy(existing_full)

    for date_key, day_payload in new_data.items():
        merged[date_key] = day_payload

    return dict(sorted(merged.items()))


def build_rolling_data(full_data: dict) -> dict:
    today = datetime.now().date()
    start_dt = today - timedelta(days=ROLLING_DAYS_BEFORE)
    end_dt = today + timedelta(days=ROLLING_DAYS_AFTER)

    rolling = {}
    for date_key, payload in full_data.items():
        try:
            dt = datetime.strptime(date_key, "%Y%m%d").date()
        except ValueError:
            log_warning(f"Skipping invalid date key in full data: {date_key}")
            continue

        if start_dt <= dt <= end_dt:
            rolling[date_key] = payload

    return dict(sorted(rolling.items()))


def validate_merged_payload(full_data: dict):
    """
    Optional sanity checks after merge.
    """
    for date_key, day_payload in full_data.items():
        if not isinstance(day_payload, dict):
            raise ValidationError(f"Date '{date_key}' payload must be an object.")

        for section in ["breakfast", "lunch", "dinner", "supper"]:
            if section not in day_payload:
                raise ValidationError(f"Date '{date_key}' missing section '{section}'.")

        # vegetarian is now expected too
        if "vegetarian" not in day_payload:
            raise ValidationError(f"Date '{date_key}' missing section 'vegetarian'.")


def main():
    os.makedirs(DATA_DIR, exist_ok=True)
    try:
        excel_paths = find_excel_files()
        existing_full = read_json_file(FULL_JSON_FILE)
        merged_full = existing_full
        warnings = []
        parsed_dates = set()

        for excel_path in excel_paths:
            excel_data, file_warnings = load_excel_data(excel_path)
            merged_full = merge_full_data(merged_full, excel_data)
            warnings.extend(file_warnings)
            parsed_dates.update(excel_data)

        validate_merged_payload(merged_full)

        rolling_data = build_rolling_data(merged_full)

        save_json_file(FULL_JSON_FILE, merged_full)
        save_json_file(ROLLING_JSON_FILE, rolling_data)

        for w in warnings:
            log_warning(w)

        log_info(f"Excel files processed: {len(excel_paths)}")
        log_info(f"Excel parsed unique dates: {len(parsed_dates)}")
        log_info(f"Full dataset dates: {len(merged_full)}")
        log_info(f"Rolling dataset dates: {len(rolling_data)}")
        log_info(f"Warnings: {len(warnings)}")
        log_info("Conversion completed successfully.")

    except ValidationError as e:
        log_error(str(e))
        sys.exit(1)
    except Exception as e:
        log_error(f"Unexpected error: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()


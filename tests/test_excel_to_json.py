import json
import os
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

import excel_to_json  # noqa: E402


WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def build_week_sheet(home_taste_rows: int):
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 1"

    for offset, weekday in enumerate(WEEKDAYS, start=3):
        ws.cell(2, offset, weekday)
        ws.cell(3, offset, date(2026, 9, 21) + timedelta(days=offset - 3))

    breakfast_start = 4
    lunch_start = 9
    main_meal_height = home_taste_rows + 5
    supper_rice_spice_rows = 5
    supper_height = supper_rice_spice_rows + 5
    dinner_start = lunch_start + main_meal_height
    supper_start = dinner_start + main_meal_height
    vegetarian_row = supper_start + supper_height

    ws.cell(breakfast_start, 1, "BREAKFAST")
    breakfast_labels = ["Bread & Spreads", "Cereal", "Main", "Protein", "Side"]
    for index, label in enumerate(breakfast_labels, start=breakfast_start):
        ws.cell(index, 2, label)

    for section_name, start_row, rice_spice_rows in (
        ("LUNCH", lunch_start, home_taste_rows),
        ("DINNER", dinner_start, home_taste_rows),
        ("SUPPER", supper_start, supper_rice_spice_rows),
    ):
        ws.cell(start_row, 1, section_name)
        ws.cell(start_row, 2, "The Bowl")
        ws.cell(start_row + 1, 2, "The Culinary Table")
        ws.cell(start_row + 2, 2, "Rice & Spice")
        ws.cell(start_row + 2 + rice_spice_rows, 2, "Green House")
        ws.cell(start_row + 3 + rice_spice_rows, 2, "Soup")
        ws.cell(start_row + 4 + rice_spice_rows, 2, "Dessert")

    ws.cell(vegetarian_row, 2, "Vegetarian")

    for col in range(3, 10):
        for row in range(breakfast_start, vegetarian_row + 1):
            ws.cell(row, col, f"R{row}C{col}")

    return ws


class DynamicLayoutTests(unittest.TestCase):
    def test_detects_september_style_six_row_home_taste_sections(self):
        ws = build_week_sheet(home_taste_rows=6)

        layout = excel_to_json.detect_sheet_layout(ws)

        self.assertEqual((layout["breakfast"]["start"], layout["breakfast"]["end"]), (4, 8))
        self.assertEqual((layout["lunch"]["start"], layout["lunch"]["end"]), (9, 19))
        self.assertEqual((layout["dinner"]["start"], layout["dinner"]["end"]), (20, 30))
        self.assertEqual((layout["supper"]["start"], layout["supper"]["end"]), (31, 40))
        self.assertEqual(layout["vegetarian_row"], 41)

    def test_detects_october_style_five_row_home_taste_sections(self):
        ws = build_week_sheet(home_taste_rows=5)

        layout = excel_to_json.detect_sheet_layout(ws)

        self.assertEqual((layout["breakfast"]["start"], layout["breakfast"]["end"]), (4, 8))
        self.assertEqual((layout["lunch"]["start"], layout["lunch"]["end"]), (9, 18))
        self.assertEqual((layout["dinner"]["start"], layout["dinner"]["end"]), (19, 28))
        self.assertEqual((layout["supper"]["start"], layout["supper"]["end"]), (29, 38))
        self.assertEqual(layout["vegetarian_row"], 39)

    def test_parses_variable_length_sections_by_row_labels(self):
        ws = build_week_sheet(home_taste_rows=5)

        parsed = excel_to_json.parse_sheet(ws, [])
        first_day = parsed["20260921"]

        self.assertEqual(len(first_day["lunch"]["homeTaste"]), 5)
        self.assertEqual(len(first_day["lunch"]["dailySpecial"]), 4)
        self.assertEqual(len(first_day["dinner"]["homeTaste"]), 5)
        self.assertEqual(len(first_day["supper"]["homeTaste"]), 6)
        self.assertEqual(len(first_day["supper"]["dailySpecial"]), 3)
        self.assertEqual(first_day["vegetarian"]["lines"], ["R39C3"])

    def test_reports_missing_structural_label(self):
        ws = build_week_sheet(home_taste_rows=5)
        ws.cell(26, 2, "")

        with self.assertRaisesRegex(excel_to_json.ValidationError, "Green House"):
            excel_to_json.detect_sheet_layout(ws)


class GithubEventSelectionTests(unittest.TestCase):
    def test_returns_all_added_and_modified_workbooks_in_event_order(self):
        event = {
            "commits": [
                {
                    "added": ["source/Hoyeh Menu Sep 26 - Updated.xlsx"],
                    "modified": ["README.md"],
                },
                {
                    "added": ["source/Hoyeh Menu Oct 26.xlsx"],
                    "modified": ["source/Hoyeh Menu Aug 26.xlsx"],
                },
            ]
        }

        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as event_file:
            json.dump(event, event_file)
            event_path = event_file.name

        try:
            with (
                patch.dict(os.environ, {"GITHUB_EVENT_PATH": event_path}),
                patch.object(
                    excel_to_json,
                    "existing_source_excel_path",
                    side_effect=lambda path: Path(path),
                ),
            ):
                selected = excel_to_json.find_excel_files_from_github_event()
        finally:
            os.unlink(event_path)

        self.assertEqual(
            [excel_to_json.normalize_repo_path(path) for path in selected],
            [
                "source/Hoyeh Menu Sep 26 - Updated.xlsx",
                "source/Hoyeh Menu Oct 26.xlsx",
                "source/Hoyeh Menu Aug 26.xlsx",
            ],
        )

    def test_repeated_path_uses_final_checkout_once_at_latest_position(self):
        event = {
            "commits": [
                {"added": ["source/Hoyeh Menu Sep 26.xlsx"]},
                {
                    "added": ["source/Hoyeh Menu Oct 26.xlsx"],
                    "modified": ["source/Hoyeh Menu Sep 26.xlsx"],
                },
            ]
        }

        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as event_file:
            json.dump(event, event_file)
            event_path = event_file.name

        try:
            with (
                patch.dict(os.environ, {"GITHUB_EVENT_PATH": event_path}),
                patch.object(
                    excel_to_json,
                    "existing_source_excel_path",
                    side_effect=lambda path: Path(path),
                ),
            ):
                selected = excel_to_json.find_excel_files_from_github_event()
        finally:
            os.unlink(event_path)

        self.assertEqual(
            [excel_to_json.normalize_repo_path(path) for path in selected],
            [
                "source/Hoyeh Menu Oct 26.xlsx",
                "source/Hoyeh Menu Sep 26.xlsx",
            ],
        )


if __name__ == "__main__":
    unittest.main()
